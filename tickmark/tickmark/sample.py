"""Read the auditor's Excel sample into SampleRow objects."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from tickmark.models import RowError, SampleLoadResult, SampleRow

ALIASES = {
    "row_id": {"row", "id", "no", "no.", "#"},
    "vendor": {"vendor", "supplier"},
    "amount": {"amount", "total"},
    "invoice_date": {"date", "invoice date"},
    "reference": {"reference", "ref"},
}
REQUIRED = ("vendor", "amount", "invoice_date")


class SampleFormatError(Exception):
    pass


def load_sample(path: str | Path, sheet: str | None = None) -> SampleLoadResult:
    path = Path(path)
    if not path.is_file():
        raise SampleFormatError(f"file not found: {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, OSError, KeyError) as e:
        raise SampleFormatError(f"not a readable .xlsx file: {path}") from e
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
    except KeyError as e:
        raise SampleFormatError(f"sheet not found: {sheet}") from e

    rows_iter = ((r[0].row, [c.value for c in r]) for r in ws.iter_rows())
    header_row, columns = None, {}
    for number, values in rows_iter:
        if any(v not in (None, "") for v in values):
            header_row, columns = number, _map_headers(values)
            break
    if header_row is None:
        raise SampleFormatError("sheet is empty")
    missing = [c for c in REQUIRED if c not in columns]
    if missing:
        raise SampleFormatError(f"missing required column(s): {', '.join(missing)}")

    rows: list[SampleRow] = []
    errors: list[RowError] = []
    for number, values in rows_iter:
        if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        try:
            rows.append(_parse_row(number, values, columns))
        except ValueError as e:
            errors.append(RowError(row_number=number, message=str(e)))
    return SampleLoadResult(rows=rows, errors=errors)


def _map_headers(values: list) -> dict[str, int]:
    columns: dict[str, int] = {}
    for idx, v in enumerate(values):
        name = str(v).strip().lower() if v is not None else ""
        for field, aliases in ALIASES.items():
            if name in aliases and field not in columns:
                columns[field] = idx
    return columns


def _parse_row(number: int, values: list, columns: dict[str, int]) -> SampleRow:
    def get(field: str):
        i = columns.get(field)
        return values[i] if i is not None and i < len(values) else None

    vendor = str(get("vendor") or "").strip()
    if not vendor:
        raise ValueError("vendor is empty")
    reference = get("reference")
    raw_id = get("row_id")
    try:
        row_id = int(raw_id) if raw_id not in (None, "") else number
    except (TypeError, ValueError):
        row_id = number
    return SampleRow(
        row_id=row_id,
        vendor=vendor,
        amount=_parse_amount(get("amount")),
        invoice_date=_parse_date(get("invoice_date")),
        reference=str(reference).strip() if reference not in (None, "") else None,
    )


def _parse_amount(value) -> Decimal:
    if value is None or isinstance(value, bool) or value == "":
        raise ValueError("amount is missing or invalid")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if "," in text and "." in text:  # last separator is the decimal one
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        text = f"{head.replace(',', '')}.{tail}" if len(tail) <= 2 else text.replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(f"amount is not a number: {value!r}") from None


def _parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise ValueError(f"date is missing or not YYYY-MM-DD: {value!r}")
