"""Read the auditor's Excel sample into SampleRow objects."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from tickmark.models import RowError, SampleLoadResult, SampleRow, SkippedRow

ALIASES = {
    "row_id": {"row", "row id", "id", "no", "#"},
    "vendor": {"vendor", "supplier", "vendor name", "supplier name"},
    "amount": {"amount", "total", "invoice amount", "invoice total"},
    "invoice_date": {"date", "invoice date"},
    "reference": {"reference", "ref", "invoice no", "invoice number", "invoice #", "invoice id"},
}
REQUIRED = ("vendor", "amount", "invoice_date")
FOOTER_WORDS = {"total", "subtotal", "sub total", "grand total", "sum", "totals", "notes", "note"}
ZERO_WIDTH = dict.fromkeys(map(ord, "​‌‍﻿"))


class SampleFormatError(Exception):
    pass


def load_sample(path: str | Path, sheet: str | None = None) -> SampleLoadResult:
    path = Path(path)
    if not path.is_file():
        raise SampleFormatError(f"file not found: {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, OSError, KeyError) as e:
        raise SampleFormatError(f"not a readable .xlsx file: {path}") from e
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
    except KeyError as e:
        raise SampleFormatError(f"sheet not found: {sheet}") from e

    rows_iter = ((r[0].row, [c.value for c in r]) for r in ws.iter_rows())
    header_row, columns = None, {}
    for number, values in rows_iter:
        if any(v not in (None, "") for v in values):
            header_row, columns = number, _map_headers(values)
            break
    if header_row is None:
        raise SampleFormatError("sheet is empty")
    missing = [c for c in REQUIRED if c not in columns]
    if missing:
        raise SampleFormatError(f"missing required column(s): {', '.join(missing)}")

    rows: list[SampleRow] = []
    errors: list[RowError] = []
    skipped: list[SkippedRow] = []
    for number, values in rows_iter:
        if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        footer = _footer_label(values)
        if footer:
            skipped.append(SkippedRow(row_number=number, reason=f"footer row: {footer!r}"))
            continue
        try:
            rows.append(_parse_row(number, values, columns))
        except ValueError as e:
            errors.append(RowError(row_number=number, message=str(e)))
    return SampleLoadResult(rows=rows, errors=errors, skipped=skipped)


def clean_text(value) -> str:
    """Remove zero-width characters, turn NBSP/tabs into spaces, collapse whitespace."""
    text = str(value or "").translate(ZERO_WIDTH).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def _footer_label(values: list) -> str | None:
    """A row whose first non-empty cell is a word like 'Total' is a footer, not data."""
    first = next((v for v in values if v not in (None, "")), None)
    if isinstance(first, str) and clean_text(first).lower().rstrip(":") in FOOTER_WORDS:
        return clean_text(first)
    return None


def _header_name(value) -> str:
    """'Amount (EUR)' -> 'amount', 'Invoice No.' -> 'invoice no', ' Row ID: ' -> 'row id'."""
    text = re.sub(r"\([^)]*\)", " ", clean_text(value).lower())
    return " ".join(text.replace(".", " ").replace(":", " ").replace("_", " ").split())


def _map_headers(values: list) -> dict[str, int]:
    columns: dict[str, int] = {}
    for idx, v in enumerate(values):
        name = _header_name(v)
        for field, aliases in ALIASES.items():
            if name in aliases and field not in columns:
                columns[field] = idx
    return columns


def _parse_row(number: int, values: list, columns: dict[str, int]) -> SampleRow:
    def get(field: str):
        i = columns.get(field)
        return values[i] if i is not None and i < len(values) else None

    vendor = clean_text(get("vendor"))
    if not vendor:
        raise ValueError("vendor is empty")
    reference = get("reference")
    return SampleRow(
        row_id=_parse_row_id(get("row_id"), number),
        vendor=vendor,
        amount=_parse_amount(get("amount")),
        invoice_date=_parse_date(get("invoice_date")),
        reference=clean_text(reference) or None,
        excel_row=number,
    )


def _parse_row_id(value, number: int) -> str:
    """Keep the auditor's id as written; 7.0 from Excel becomes '7'. Empty -> the Excel row number."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value) or str(number)


def _parse_amount(value) -> Decimal:
    if value is None or isinstance(value, bool) or value == "":
        raise ValueError("amount is missing or invalid")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = clean_text(value)
    negative = False
    if raw.startswith("(") and raw.endswith(")"):  # accounting negative: (1,234.50)
        negative, raw = True, raw[1:-1]
    elif raw.endswith("-"):  # trailing minus: 123.45-
        negative, raw = True, raw[:-1]
    text = re.sub(r"[^\d,.\-]", "", raw)
    if "," in text and "." in text:  # last separator is the decimal one
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        text = f"{head.replace(',', '')}.{tail}" if len(tail) <= 2 else text.replace(",", "")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"amount is not a number: {value!r}") from None
    return -amount if negative else amount


def _parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        try:
            return date.fromisoformat(text)
        except ValueError:
            pass
        if m := re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", text):  # dd.mm.yyyy
            try:
                return date(int(m[3]), int(m[2]), int(m[1]))
            except ValueError:
                pass
        if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", text):
            raise ValueError(f"ambiguous date (day/month order unknown): {value!r}")
    raise ValueError(f"date is missing or not YYYY-MM-DD / DD.MM.YYYY: {value!r}")
