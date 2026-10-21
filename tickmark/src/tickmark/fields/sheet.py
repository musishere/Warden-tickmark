"""Read a worksheet without any schema: the header row and every row's cells, as received."""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from tickmark.config import settings
from tickmark.definitions import DEFINITIONS
from tickmark.fields.types import clean_text
from tickmark.log import get
from tickmark.models import Cell

log = get("sheet")


class SampleFormatError(Exception):
    pass


@dataclass
class SheetRow:
    excel_row: int
    values: list  # Python values as openpyxl read them (datetime, float, str, None)


@dataclass
class Sheet:
    title: str
    header_row: int
    headers: list[Cell]  # one per column: letter + header text as written
    rows: list[SheetRow]  # every row below the header, blank ones included
    preamble: list[SheetRow] = field(default_factory=list)  # non-blank rows above the header (titles, notes)
    formulas: set[tuple[int, str]] = field(default_factory=set)  # (excel row, column) of formulas never calculated

    def cells(self, row: SheetRow) -> list[Cell]:
        """The row as JSON-safe cells, each with its column letter and header."""
        return [Cell(col=h.col, header=h.header, value=json_value(v), kind=kind(v))
                for h, v in zip(self.headers, padded(row.values, len(self.headers)), strict=True)]


CODE_IN_BRACKETS = re.compile(r"\(\s*([A-Z]{3})\s*\)")  # 'Amount (EUR)': a currency code, as written in capitals


def header_name(value) -> str:
    """'Amount (EUR)' -> 'amount', 'Invoice No.' -> 'invoice no', ' Row ID: ' -> 'row id'.
    Only a currency code in brackets is dropped: other bracketed words change the meaning and stay,
    so 'Amount (excl. VAT)' -> 'amount excl vat' and 'Date (posting)' -> 'date posting' match no alias."""
    text = CODE_IN_BRACKETS.sub(" ", clean_text(value)).lower()
    for ch in ".:_()":
        text = text.replace(ch, " ")
    return " ".join(text.split())


KNOWN_HEADERS = {a for d in DEFINITIONS.values() for f in d.fields for a in f.aliases}


def header_score(values: list) -> tuple[int, int]:
    """How much a row looks like a header: known header names first, then text cells."""
    texts = [v for v in values if isinstance(v, str) and clean_text(v)]
    return sum(header_name(v) in KNOWN_HEADERS for v in texts), len(texts)


def padded(values: list, width: int) -> list:
    return list(values) + [None] * (width - len(values))


def kind(v) -> str:
    if v is None or v == "":
        return "empty"
    return type(v).__name__


def json_value(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    return v


def is_blank(values: list) -> bool:
    return all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in values)


def read_sheet(path: str | Path, sheet: str | None = None) -> Sheet:
    """The header is the most header-like of the first settings.header_scan_rows non-blank rows
    (header_score; the earliest on a tie), so a title block above it is kept, not read as headers."""
    path = Path(path)
    if not path.is_file():
        raise SampleFormatError(f"file not found: {path}")
    try:
        wb = load_workbook(path, data_only=True)
        # ponytail: second load only to tell "formula never calculated" from "empty"; fine at sample sizes
        wb_formulas = load_workbook(path)
    except (InvalidFileException, OSError, KeyError) as e:
        raise SampleFormatError(f"not a readable .xlsx file: {path}") from e
    try:
        ws, ws_f = (wb[sheet], wb_formulas[sheet]) if sheet else (wb.worksheets[0], wb_formulas.worksheets[0])
    except KeyError as e:
        raise SampleFormatError(f"sheet not found: {sheet}") from e

    all_rows = [(r[0].row, [c.value for c in r]) for r in ws.iter_rows()]
    filled = [i for i, (_, values) in enumerate(all_rows) if not is_blank(values)]
    if not filled:
        raise SampleFormatError("sheet is empty")
    first = max(filled[:settings.header_scan_rows], key=lambda i: header_score(all_rows[i][1]))
    header_row, header_values = all_rows[first]
    width = max(len(values) for _, values in all_rows)
    headers = [Cell(col=get_column_letter(i + 1), header=clean_text(v) or None,
                    value=json_value(v), kind=kind(v))
               for i, v in enumerate(padded(header_values, width))]
    formulas = {(c.row, c.column_letter) for r in ws_f.iter_rows(min_row=header_row + 1) for c in r
                if c.data_type == "f" and ws[c.coordinate].value is None}
    log.info("%s, sheet %r: header on row %d (%d known names, %d text cells), %d columns, %d rows below",
             path.name, ws.title, header_row, *header_score(header_values), len(headers), len(all_rows) - first - 1)
    if above := [all_rows[i][0] for i in filled if i < first]:
        log.info("rows above the header kept as skipped: %s", above)
    if formulas:
        log.warning("%d formula cells were never calculated (file not saved in Excel)", len(formulas))
    return Sheet(ws.title, header_row, headers, [SheetRow(n, values) for n, values in all_rows[first + 1:]],
                 preamble=[SheetRow(*all_rows[i]) for i in filled if i < first], formulas=formulas)
