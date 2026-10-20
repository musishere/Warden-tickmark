"""Write one identification run to the audit database, completely or not at all, inside one
tenant + engagement (see tickmark.store.context)."""
import hashlib
import json
import subprocess
from datetime import UTC, date, datetime
from decimal import Decimal
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import Connection, Engine, select
from sqlalchemy.dialects.postgresql import insert

from tickmark.match.identify import THRESHOLDS
from tickmark.models import ExtractedInvoice, Identification
from tickmark.pipeline import IdentificationRun
from tickmark.store import evidence
from tickmark.store.context import TenantContext, scoped
from tickmark.store.schema import (
    documents,
    extractions,
    match_results,
    run_documents,
    runs,
    sample_issues,
    sample_rows,
)


def matcher_version() -> str:
    try:
        pkg = version("tickmark")
    except PackageNotFoundError:
        pkg = "unknown"
    here = Path(__file__).resolve().parent
    try:
        commit = subprocess.run(["git", "rev-parse", "--short=12", "HEAD"], cwd=here, capture_output=True,
                                text=True, check=True).stdout.strip()
        dirty = subprocess.run(["git", "status", "--porcelain", "--", str(here.parents[1])], cwd=here,
                               capture_output=True, text=True, check=True).stdout.strip()
        return f"tickmark {pkg} @ {commit}{'+dirty' if dirty else ''}"
    except (OSError, subprocess.CalledProcessError):
        return f"tickmark {pkg} @ no-git"


def _json_cell(v):
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    return v


def raw_cells_by_row(sample: Path) -> tuple[str, dict[int, list]]:
    """The first sheet's cells exactly as the loader saw them, per Excel row."""
    ws = load_workbook(sample, data_only=True).worksheets[0]
    return ws.title, {r[0].row: [_json_cell(c.value) for c in r] for r in ws.iter_rows()}


def canonical_sha256(data) -> str:
    return hashlib.sha256(json.dumps(data, sort_keys=True, separators=(",", ":"), default=str)
                          .encode()).hexdigest()


def _scope(ctx: TenantContext) -> dict:
    return {"tenant_id": ctx.tenant_id, "engagement_id": ctx.engagement_id}


def _document(conn: Connection, ctx: TenantContext, path: Path, kind: str, pages: int | None = None) -> UUID:
    """The file, by hash, within this engagement (never shared with another engagement)."""
    sha, stored = evidence.put(path, ctx.tenant_id, ctx.engagement_id)
    conn.execute(insert(documents).values(**_scope(ctx), sha256=sha, kind=kind, size_bytes=Path(path).stat().st_size,
                                          pages=pages, storage_path=str(stored))
                 .on_conflict_do_nothing(constraint="documents_once_per_engagement"))
    return conn.execute(select(documents.c.id).where(
        documents.c.tenant_id == ctx.tenant_id, documents.c.engagement_id == ctx.engagement_id,
        documents.c.sha256 == sha)).scalar_one()


def _extraction(conn: Connection, ctx: TenantContext, document_id: UUID, inv: ExtractedInvoice) -> UUID:
    raw = inv.raw if inv.raw is not None else {}
    values = {**_scope(ctx), "document_id": document_id, "extractor": inv.extractor or "unknown",
              "model_id": inv.model_id, "api_version": inv.api_version, "raw_sha256": canonical_sha256(raw),
              "raw": raw, "fields": {k: f.model_dump(mode="json") for k, f in inv.fields.items()}}
    conn.execute(insert(extractions).values(**values).on_conflict_do_nothing(constraint="extractions_once"))
    return conn.execute(select(extractions.c.id).where(
        extractions.c.document_id == document_id, extractions.c.extractor == values["extractor"],
        extractions.c.model_id == values["model_id"], extractions.c.raw_sha256 == values["raw_sha256"],
    )).scalar_one()


def _pages(inv: ExtractedInvoice) -> int | None:
    raw = inv.raw or {}
    pages = raw.get("analyzeResult", raw).get("pages") if inv.extractor == "azure-docint" else None
    return len(pages) if pages else None


def _result_values(r: Identification, run_id: UUID, sample_row_id: UUID, doc_ids: dict[str, UUID]) -> dict:
    return {
        "run_id": run_id, "sample_row_id": sample_row_id, "status": r.status,
        "document_id": doc_ids[r.invoice_file] if r.invoice_file else None,
        "page": r.snips[0].page if r.snips else None,
        "evidence": {"invoice_file": r.invoice_file, "lines": r.evidence,
                     "snips": [s.model_dump(mode="json") for s in r.snips]},
        "candidates": [{**c.model_dump(mode="json"), "document_id": str(doc_ids[c.invoice_file])}
                       for c in r.candidates],
        "reason": r.reason,
        "comparisons": [c.model_dump(mode="json") for c in r.comparisons],
        "steps": r.steps,
    }


def save_run(run: IdentificationRun, ctx: TenantContext, *, started_at: datetime, db: Engine | None = None) -> UUID:
    """Store a finished run in ctx's engagement: documents (by hash), extractions, rows, issues,
    results. One transaction: all of it or nothing."""
    sheet, raw = raw_cells_by_row(run.sample_path)
    model_ids = sorted({inv.model_id for inv in run.extracted if inv.model_id})
    scope = _scope(ctx)
    with scoped(ctx, db) as conn:
        sample_doc = _document(conn, ctx, run.sample_path, "sample")
        run_id = conn.execute(runs.insert().values(
            **scope, run_by=ctx.user, started_at=started_at,
            finished_at=datetime.now(UTC), status="completed", sample_document_id=sample_doc,
            sheet=sheet, matcher_version=matcher_version(), thresholds=THRESHOLDS,
            extractor=",".join(sorted({inv.extractor for inv in run.extracted})) or "none",
            model_id=",".join(model_ids) or None,
        ).returning(runs.c.id)).scalar_one()

        doc_ids: dict[str, UUID] = {}
        for path, inv in zip(run.invoice_files, run.extracted, strict=True):
            doc_id = _document(conn, ctx, path, "pdf", _pages(inv))
            doc_ids[inv.file] = doc_id
            conn.execute(run_documents.insert().values(**scope, run_id=run_id, document_id=doc_id, file_name=inv.file,
                                                       extraction_id=_extraction(conn, ctx, doc_id, inv)))

        row_ids: dict[int, UUID] = {}
        flags = {cr.row.excel_row: cr for cr in run.cleaned.rows}
        for row in run.loaded.rows:
            cr = flags[row.excel_row]
            row_ids[row.excel_row] = conn.execute(sample_rows.insert().values(
                **scope, run_id=run_id, row_id=row.row_id, excel_row=row.excel_row, vendor=row.vendor,
                amount=row.amount, invoice_date=row.invoice_date, reference=row.reference,
                vendor_key=cr.vendor_key, flags=cr.flags, raw_cells=raw.get(row.excel_row, []),
            ).returning(sample_rows.c.id)).scalar_one()
        issues = [(e.row_number, "error", e.message) for e in run.loaded.errors]
        issues += [(s.row_number, "skipped", s.reason) for s in run.loaded.skipped]
        for n, kind, reason in issues:
            conn.execute(sample_issues.insert().values(**scope, run_id=run_id, excel_row=n, kind=kind, reason=reason,
                                                       raw_cells=raw.get(n, [])))

        for r in run.results:
            conn.execute(match_results.insert().values(
                **scope, **_result_values(r, run_id, row_ids[r.excel_row], doc_ids)))
    return run_id


def save_failed_run(ctx: TenantContext, *, sample_path: Path | None, started_at: datetime, error: str,
                    extractor: str = "unknown", db: Engine | None = None) -> UUID:
    """A run that crashed is still recorded, with what is known about it."""
    with scoped(ctx, db) as conn:
        sample_doc = (_document(conn, ctx, sample_path, "sample")
                      if sample_path and Path(sample_path).is_file() else None)
        return conn.execute(runs.insert().values(
            **_scope(ctx), run_by=ctx.user, started_at=started_at,
            finished_at=datetime.now(UTC), status="failed", error=error,
            sample_document_id=sample_doc, matcher_version=matcher_version(), thresholds=THRESHOLDS,
            extractor=extractor,
        ).returning(runs.c.id)).scalar_one()
