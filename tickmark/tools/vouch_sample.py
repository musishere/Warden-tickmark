"""Check each sample row's claim against its invoice PDF: ✅ matched · ⚠️ difference · 👀 needs review · ❓ no evidence.

    python -m tools.vouch_sample data/external/tickmark-sample/sample.xlsx data/external/tickmark-sample/invoices \
        --labels data/external/tickmark-sample/labels --period 2026-01-01:2026-12-31 [--trace] [--out x.xlsx]
    python -m tools.vouch_sample SAMPLE.xlsx INVOICE_DIR --save      # period from the engagement

First finds which PDF each row is about (tools.identify_sample), then compares the row's vendor,
amount, date and invoice number with that PDF and runs the invoice's own checks (totals add up,
lines sum, VAT rate, inside the audit period, duplicates). --trace shows every comparison and check;
--out writes them to a workbook (Verdicts, Agreement, Checks + the identification sheets).
"""
import argparse
from datetime import UTC, datetime
from pathlib import Path

from openpyxl.styles import Alignment, PatternFill

from tickmark.extract import DocIntExtractor, LabelExtractor
from tickmark.models import Verdict
from tickmark.pipeline import run_vouch
from tickmark.vouch import MARK, VouchPolicy, parse_period
from tools.identify_sample import build_workbook, pdf_linker, save_workbook, style
from tools.tenant_args import add_tenant_args, tenant_context

RESULT_MARK = {"agrees": "✓", "differs": "✗", "unverifiable": "?", "pass": "✓", "fail": "✗", "warn": "!",
               "not_checkable": "?"}
FILL = {"matched": "E2F0D9", "difference": "FCE4D6", "needs_review": "FFF2CC", "no_evidence": "F2F2F2",
        "agrees": "E2F0D9", "pass": "E2F0D9", "differs": "FCE4D6", "fail": "FCE4D6", "warn": "FFF2CC",
        "unverifiable": "FFF2CC", "not_checkable": "FFF2CC"}


def print_verdicts(verdicts: list[Verdict], trace: bool) -> None:
    for v in verdicts:
        print(f"{MARK[v.verdict]} {v.row_id:>5}  {v.verdict:<13} {v.invoice_file or ''}")
        for r in v.reasons:
            print(f"{'':21}· {r}")
        if trace and v.agreements:
            print(f"{'':6}Row vs PDF:")
            for a in v.agreements:
                print(f"{'':8}{RESULT_MARK[a.result]} {a.field:<7} row {a.row_value:<28} pdf {a.pdf_value or '—':<28} "
                      f"{a.result}{'  ' + (a.delta or a.detail) if (a.delta or a.detail) else ''}")
            print(f"{'':6}Invoice checks:")
            for c in v.checks:
                print(f"{'':8}{RESULT_MARK[c.status]} {c.name:<17} {c.status:<14} {c.detail}")
        if trace:
            print(f"{'':6}Steps:")
            for i, s in enumerate(v.steps, 1):
                print(f"{'':8}{i}. {s}")
            print()
    counts = {k: sum(v.verdict == k for v in verdicts) for k in MARK}
    print(f"\n{len(verdicts)} rows | ✅ matched {counts['matched']} | ⚠️ difference {counts['difference']} | "
          f"👀 needs review {counts['needs_review']} | ❓ no evidence {counts['no_evidence']}")


def add_verdict_sheets(wb, verdicts: list[Verdict], pdfs) -> None:
    link = pdf_linker(pdfs)
    ws = wb.create_sheet("Verdicts", 0)
    ws.append(["Excel row", "Row", "Verdict", "Invoice PDF", "Why (each difference, failed check, warning)"])
    for v in verdicts:
        ws.append([v.excel_row, v.row_id, f"{MARK[v.verdict]} {v.verdict}", v.invoice_file, "\n".join(v.reasons)])
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=FILL[v.verdict])
        ws.cell(ws.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
        link(ws.cell(ws.max_row, 4), v.invoice_file)

    ag = wb.create_sheet("Agreement", 1)
    ag.append(["Row", "PDF", "Field", "Row value", "PDF value", "Result", "Difference", "Detail"])
    for v in verdicts:
        for a in v.agreements:
            ag.append([v.row_id, v.invoice_file, a.field, a.row_value, a.pdf_value, a.result, a.delta, a.detail])
            ag.cell(ag.max_row, 6).fill = PatternFill("solid", fgColor=FILL[a.result])
            link(ag.cell(ag.max_row, 2), v.invoice_file)
    ag.auto_filter.ref = ag.dimensions

    ck = wb.create_sheet("Checks", 2)
    ck.append(["Row", "PDF", "Check", "Status", "Detail"])
    for v in verdicts:
        for c in v.checks:
            ck.append([v.row_id, v.invoice_file, c.name, c.status, c.detail])
            ck.cell(ck.max_row, 4).fill = PatternFill("solid", fgColor=FILL[c.status])
            link(ck.cell(ck.max_row, 2), v.invoice_file)
    ck.auto_filter.ref = ck.dimensions
    for sheet, widths in ((ws, [10, 8, 18, 22, 110]), (ag, [8, 22, 8, 30, 30, 13, 22, 40]),
                          (ck, [8, 22, 18, 14, 100])):
        style(sheet, widths)
    wb.active = 0


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("sample", type=Path)
    ap.add_argument("invoices", type=Path)
    ap.add_argument("--labels", type=Path, help="use label JSON instead of Azure (testing only)")
    ap.add_argument("--period", help="audit period YYYY-MM-DD:YYYY-MM-DD (default with --save: the engagement's)")
    ap.add_argument("--trace", action="store_true", help="print every comparison, check and step")
    ap.add_argument("--out", type=Path, help="write verdicts, comparisons and checks to a workbook")
    ap.add_argument("--save", action="store_true", help="store the run (and verdicts) in the audit database")
    add_tenant_args(ap)
    args = ap.parse_args()
    ctx = tenant_context(ap, args) if args.save else None

    period = parse_period(args.period) if args.period else None
    if period is None and ctx:
        from tickmark.store import engagement_period

        period = engagement_period(ctx)
    if not period or None in period:
        print("⚠ no audit period: in_period cannot be checked, so no row can be ✅ "
              "(pass --period, or set it: python -m tools.tenants set-period)\n")
        period = (None, None)
    policy = VouchPolicy(period_start=period[0], period_end=period[1])

    started = datetime.now(UTC)
    extractor = LabelExtractor(args.labels) if args.labels else DocIntExtractor()
    run = run_vouch(args.sample, args.invoices, extractor, policy)
    print_verdicts(run.verdicts, args.trace)
    if args.out:
        wb = build_workbook(run.results, run.extracted, args.invoices)
        add_verdict_sheets(wb, run.verdicts, args.invoices)
        save_workbook(wb, args.out)
        print(f"wrote {args.out}")
    if args.save:
        from tickmark.store import save_run

        run_id = save_run(run, ctx, started_at=started)
        scope = f"--tenant {ctx.tenant_id} --engagement {ctx.engagement_id}"
        print(f"saved run {run_id}\n  show:   python -m tools.show_run {run_id} {scope} --trace")


if __name__ == "__main__":
    main()
