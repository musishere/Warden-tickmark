"""Measure how accurately sample rows are linked to their invoice PDFs.

    python -m tools.eval_identify            # Azure Document Intelligence (cached)
    python -m tools.eval_identify --labels   # perfect extraction from labels: tests matching alone

Per dataset: wrong links (the number that must be 0), correct links, rows sent to review or left
without evidence, missed links, and whether each row got its safe outcome. Exits 1 on any wrong link.
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from tools.identify_sample import run

GEN, EXT, HARD = Path("data/generated"), Path("data/external/tickmark-sample"), Path("data/generated/hard")


def truth_generated() -> dict[str, dict]:
    labels = json.loads((GEN / "labels" / "sample_labels.json").read_text())
    return {str(t["row_id"]): {"invoice_file": t["invoice_file"],
                               "expect": "link" if t["invoice_file"] else "none"} for t in labels}


def truth_external() -> dict[str, dict]:
    """The pack's sample names the invoice number; the PDF is named after it, if it exists."""
    ws = load_workbook(EXT / "sample.xlsx", data_only=True).worksheets[0]
    truth = {}
    for row_id, _, number, *_ in ws.iter_rows(min_row=2, values_only=True):
        if row_id and number:
            f = f"{number}.pdf"
            exists = (EXT / "invoices" / f).exists()
            truth[row_id] = {"invoice_file": f if exists else None, "expect": "link" if exists else "none"}
    return truth


def truth_hard() -> dict[str, dict]:
    return {t["row_id"]: t for t in json.loads((HARD / "truth.json").read_text())}


DATASETS = {
    "generated": (GEN / "sample.xlsx", GEN / "invoices", GEN / "labels", truth_generated),
    "external": (EXT / "sample.xlsx", EXT / "invoices", EXT / "labels", truth_external),
    "hard": (HARD / "sample.xlsx", HARD / "invoices", HARD / "labels", truth_hard),
}
OUTCOME = {"linked": "link", "needs_review": "review", "no_evidence": "none"}


def evaluate(name: str, use_labels: bool) -> int:
    sample, invoices, labels, truth_fn = DATASETS[name]
    if not sample.exists():
        print(f"{name}: skipped ({sample} missing)\n")
        return 0
    truth = truth_fn()
    results, _ = run(sample, invoices, labels if use_labels else None)
    wrong = correct = missed = safe = 0
    print(f"== {name} ({len(results)} rows)")
    for r in results:
        t = truth.get(r.row_id, {"invoice_file": None, "expect": "none"})
        linked_wrong = r.status == "linked" and r.invoice_file != t["invoice_file"]
        wrong += linked_wrong
        correct += r.status == "linked" and not linked_wrong
        missed += r.status != "linked" and t["invoice_file"] is not None
        ok = OUTCOME[r.status] == t["expect"] and not linked_wrong
        safe += ok
        if not ok:
            print(f"   {'❌ WRONG LINK' if linked_wrong else '·  differs'} {r.row_id}: got {r.status} "
                  f"{r.invoice_file or ''} | expected {t['expect']} {t['invoice_file'] or ''} "
                  f"{r.reason}")
    n = len(results)
    print(f"   wrong links {wrong} | correct links {correct} | review "
          f"{sum(r.status == 'needs_review' for r in results)} | no evidence "
          f"{sum(r.status == 'no_evidence' for r in results)} | missed {missed}")
    print(f"   safe outcome {safe}/{n}   link precision "
          f"{correct / (correct + wrong):.0%}" if correct + wrong else f"   safe outcome {safe}/{n}")
    print()
    return wrong


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--labels", action="store_true", help="extract from label JSON instead of Azure")
    ap.add_argument("datasets", nargs="*", choices=list(DATASETS), help="default: all")
    args = ap.parse_args()
    print(f"extraction: {'labels (perfect, tests matching only)' if args.labels else 'Azure Document Intelligence'}\n")
    wrong = sum(evaluate(d, args.labels) for d in args.datasets or DATASETS)
    print("✅ no wrong links" if not wrong else f"❌ {wrong} wrong link(s)")
    sys.exit(1 if wrong else 0)


if __name__ == "__main__":
    main()
