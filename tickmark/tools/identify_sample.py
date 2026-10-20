"""Show which invoice PDF each sample row is about, with the evidence for every link.

    python -m tools.identify_sample data/external/tickmark-sample/sample.xlsx \
        data/external/tickmark-sample/invoices [--trace] [--out data/output/identified.xlsx]
    python -m tools.identify_sample data/generated/sample.xlsx data/generated/invoices \
        --labels data/generated/labels      # perfect extraction from labels, no Azure

Without --labels, PDFs are read with Azure Document Intelligence (cached in data/cache/docint).

--trace prints the whole process: what was read from each PDF (and what is not trusted, and why),
each row compared field by field with every PDF that agrees on anything, and the decision steps.
--out writes the same as a workbook to check by hand, with links that open each PDF:
  Identified   one line per row: status, PDF, evidence or reason
  Steps        the decision for each row, step by step
  Comparisons  every row x every PDF, field by field (row value | PDF value | result)
  PDF fields   everything read from each PDF, confidence, trusted or why not, page + position
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from tickmark.extract import DocIntExtractor, LabelExtractor, extract_all
from tickmark.match import identify
from tickmark.match.identify import pdf_fields
from tickmark.models import ExtractedInvoice, Identification
from tickmark.sample import clean_sample, load_sample

MARK = {"linked": "🔗", "needs_review": "👀", "no_evidence": "❓"}
RESULT_MARK = {"equal": "✓", "similar": "≈", "differs": "✗", "untrusted": "⚠", "missing": "∅",
               "not in row": "·"}
FILL = {"linked": "E2F0D9", "needs_review": "FFF2CC", "no_evidence": "F2F2F2",
        "equal": "E2F0D9", "similar": "FFF2CC", "differs": "FCE4D6", "untrusted": "FFF2CC"}


def run(sample: Path, invoices: Path, labels: Path | None
        ) -> tuple[list[Identification], list[ExtractedInvoice]]:
    extractor = LabelExtractor(labels) if labels else DocIntExtractor()
    rows = clean_sample(load_sample(sample).rows).rows
    extracted = extract_all(extractor, invoices)
    return identify(rows, extracted), extracted


def print_extraction(extracted: list[ExtractedInvoice]) -> None:
    print("── What was read from each PDF " + "─" * 50)
    for inv in extracted:
        print(f"{inv.file}")
        for f in pdf_fields(inv):
            conf = f"{f.confidence:.2f}" if f.confidence is not None else "  — "
            where = f"  p{f.snip.page} @ ({f.snip.bbox.x0:.0f},{f.snip.bbox.y0:.0f})" if f.snip else ""
            trust = "" if f.trusted else f"   ⚠ not used: {f.why_not}"
            print(f"    {f.name:<15} {f.value:<32} conf {conf}{where}{trust}")
    print()


def print_trace(r: Identification) -> None:
    print("    Steps:")
    for i, s in enumerate(r.steps, 1):
        print(f"      {i}. " + s.replace("\n", "\n         "))
    shown = [c for c in r.comparisons if any(k.result in ("equal", "similar") for k in c.checks)]
    if shown:
        print("    Field by field (PDFs that agree on anything):")
    for c in shown:
        print(f"      {c.invoice_file}  → {c.outcome}")
        for k in c.checks:
            print(f"        {RESULT_MARK[k.result]} {k.field:<7} row {k.row_value or '—':<30} "
                  f"pdf {k.pdf_value or '—':<30} {k.result}{'  ' + k.detail if k.detail else ''}")
    rest = len(r.comparisons) - len(shown)
    if rest:
        print(f"      ({rest} other PDF(s) agree on no field; see the Comparisons sheet)")


def write_workbook(out: Path, results: list[Identification], extracted: list[ExtractedInvoice],
                   invoice_dir: Path) -> None:
    def link(cell, file: str | None) -> None:
        if file:
            cell.hyperlink = (invoice_dir / file).resolve().as_uri()
            cell.font = Font(color="0563C1", underline="single")

    wb = Workbook()
    ws = wb.active
    ws.title = "Identified"
    ws.append(["Excel row", "Row", "Status", "Invoice PDF", "Evidence / reason", "Candidates"])
    for r in results:
        ws.append([r.excel_row, r.row_id, r.status, r.invoice_file, "; ".join(r.evidence) or r.reason,
                   "; ".join(f"{c.invoice_file} ({', '.join(c.agree)})" for c in r.candidates)])
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=FILL[r.status])
        link(ws.cell(ws.max_row, 4), r.invoice_file)

    steps = wb.create_sheet("Steps")
    steps.append(["Row", "Status", "#", "Step"])
    for r in results:
        for i, s in enumerate(r.steps, 1):
            steps.append([r.row_id, r.status, i, s])
            steps.cell(steps.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")

    comp = wb.create_sheet("Comparisons")
    comp.append(["Row", "PDF", "Outcome", "Field", "Row value", "PDF value", "Result", "Detail"])
    for r in results:
        for c in r.comparisons:
            for k in c.checks:
                comp.append([r.row_id, c.invoice_file, c.outcome, k.field, k.row_value, k.pdf_value,
                             k.result, k.detail])
                link(comp.cell(comp.max_row, 2), c.invoice_file)
                if k.result in FILL:
                    comp.cell(comp.max_row, 7).fill = PatternFill("solid", fgColor=FILL[k.result])
    comp.auto_filter.ref = comp.dimensions

    pf = wb.create_sheet("PDF fields")
    pf.append(["PDF", "Field", "Value", "Confidence", "Used as evidence", "Why not", "Page",
               "Text on page", "x0", "y0", "x1", "y1"])
    for inv in extracted:
        for f in pdf_fields(inv):
            s = f.snip
            pf.append([inv.file, f.name, f.value, f.confidence, "yes" if f.trusted else "NO", f.why_not,
                       s.page if s else None, s.text if s else None,
                       *((round(s.bbox.x0), round(s.bbox.y0), round(s.bbox.x1), round(s.bbox.y1)) if s else ())])
            link(pf.cell(pf.max_row, 1), inv.file)
            if not f.trusted:
                pf.cell(pf.max_row, 5).fill = PatternFill("solid", fgColor=FILL["untrusted"])
    pf.auto_filter.ref = pf.dimensions

    widths = {ws: [10, 8, 14, 22, 70, 50], steps: [8, 14, 4, 110], comp: [8, 22, 50, 8, 30, 30, 11, 40],
              pf: [22, 15, 32, 11, 16, 45, 6, 24, 6, 6, 6, 6]}
    for sheet, ws_widths in widths.items():
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col, w in zip("ABCDEFGHIJKL", ws_widths):
            sheet.column_dimensions[col].width = w
        sheet.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("sample", type=Path)
    ap.add_argument("invoices", type=Path)
    ap.add_argument("--labels", type=Path, help="use label JSON instead of Azure (testing only)")
    ap.add_argument("--trace", action="store_true", help="print every step of the process")
    ap.add_argument("--out", type=Path, help="write the results and the full trace to a workbook")
    args = ap.parse_args()
    results, extracted = run(args.sample, args.invoices, args.labels)

    if args.trace:
        print_extraction(extracted)
        print("── Each row " + "─" * 69)
    for r in results:
        print(f"{MARK[r.status]} {r.row_id:>5}  {r.status:<12} {r.invoice_file or ''}")
        for e in r.evidence:
            print(f"{'':20}· {e}")
        if r.reason:
            print(f"{'':20}{r.reason}")
            if not args.trace:
                for c in r.candidates:
                    print(f"{'':20}- {c.invoice_file}: agrees on {', '.join(c.agree) or 'nothing'}"
                          + (f" ({c.note})" if c.note else ""))
        if args.trace:
            print_trace(r)
            print()
    counts = {s: sum(r.status == s for r in results) for s in MARK}
    print(f"\n{len(results)} rows | linked {counts['linked']} | needs review {counts['needs_review']}"
          f" | no evidence {counts['no_evidence']}")

    if args.out:
        write_workbook(args.out, results, extracted, args.invoices)
        print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
