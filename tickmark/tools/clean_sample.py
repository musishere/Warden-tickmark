"""Load + clean an auditor sample and write the result to a new workbook.

    python -m tools.clean_sample data/generated/sample_messy.xlsx data/output/sample_messy_cleaned.xlsx

Sheet "Clean": one row per loaded sample row, with its source Excel row, vendor key and flags.
Sheet "Issues": rows that were not loaded (errors and skipped footers), with the reason.
"""
import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from tickmark.sample import clean_sample, load_sample


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("dst", type=Path)
    ap.add_argument("--period", help="YYYY-MM-DD:YYYY-MM-DD")
    args = ap.parse_args()
    period = tuple(date.fromisoformat(d) for d in args.period.split(":")) if args.period else None

    loaded = load_sample(args.src)
    cleaned = clean_sample(loaded.rows, period)

    wb = Workbook()
    ws = wb.active
    ws.title = "Clean"
    ws.append(["Excel row", "Row", "Vendor", "Vendor key", "Amount", "Date", "Flags"])
    for cr in cleaned.rows:
        r = cr.row
        ws.append([r.excel_row, r.row_id, r.vendor, cr.vendor_key, r.amount, r.invoice_date,
                   ", ".join(cr.flags)])
        ws.cell(ws.max_row, 5).number_format = "#,##0.00"
        ws.cell(ws.max_row, 6).number_format = "yyyy-mm-dd"

    issues = wb.create_sheet("Issues")
    issues.append(["Excel row", "Outcome", "Reason"])
    problems = [(e.row_number, "error", e.message) for e in loaded.errors]
    problems += [(s.row_number, "skipped", s.reason) for s in loaded.skipped]
    for p in sorted(problems):
        issues.append(list(p))

    for sheet, widths in ((ws, [10, 6, 32, 28, 12, 12, 24]), (issues, [10, 10, 60])):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col, w in zip("ABCDEFG", widths):
            sheet.column_dimensions[col].width = w
        sheet.freeze_panes = "A2"

    wb.save(args.dst)
    print(f"{args.src} -> {args.dst}: {len(cleaned.rows)} clean rows, {len(problems)} issues")


if __name__ == "__main__":
    main()
