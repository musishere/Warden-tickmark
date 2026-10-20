"""Show raw Excel cells next to what the loader and cleaner produced, row by row.

    python -m tools.compare_sample data/generated/sample_messy.xlsx [--labels data/generated/labels/sample_messy_labels.json]
    python -m tools.compare_sample data/generated/sample.xlsx --period 2025-01-01:2025-12-31
"""
import argparse
import json
from decimal import Decimal
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from tickmark.sample import clean_sample, load_sample


def actual_by_row(path: Path, period):
    loaded = load_sample(path)
    cleaned = clean_sample(loaded.rows, period)
    out: dict[int, dict] = {}
    for cr in cleaned.rows:
        r = cr.row
        out[r.excel_row] = {
            "outcome": "loaded", "vendor": r.vendor, "vendor_key": cr.vendor_key,
            "amount": str(r.amount), "date": r.invoice_date.isoformat(), "flags": cr.flags,
        }
    for e in loaded.errors:
        out[e.row_number] = {"outcome": "error", "reason": e.message}
    for s in loaded.skipped:
        out[s.row_number] = {"outcome": "skipped", "reason": s.reason}
    return out, cleaned


def agrees(exp: dict, act: dict) -> bool:
    if exp["outcome"] != act["outcome"]:
        return False
    if exp["outcome"] == "loaded":
        return (exp["vendor_key"], exp["flags"]) == (act["vendor_key"], act["flags"]) and Decimal(
            exp["amount"]) == Decimal(act["amount"])
    if "reason" in exp:
        return exp["reason"] in act["reason"]
    return True


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--labels", type=Path)
    ap.add_argument("--period", help="YYYY-MM-DD:YYYY-MM-DD")
    args = ap.parse_args()
    period = tuple(date.fromisoformat(d) for d in args.period.split(":")) if args.period else None
    expected = {e["excel_row"]: e for e in json.loads(args.labels.read_text())} if args.labels else {}

    ws = load_workbook(args.xlsx, data_only=True).worksheets[0]
    actual, cleaned = actual_by_row(args.xlsx, period)

    print(f"File: {args.xlsx}   Sheet: {ws.title!r}\n")
    ok = bad = 0
    for row in ws.iter_rows(min_row=2):
        n = row[0].row
        raw = [c.value for c in row]
        raw_txt = " | ".join(repr(v.date() if hasattr(v, "date") and not isinstance(v, date) else v)
                             if v is not None else "·" for v in raw)
        act = actual.get(n, {"outcome": "blank"})
        mark = ""
        if expected:
            good = agrees(expected[n], act)
            ok, bad = ok + good, bad + (not good)
            mark = "✅ " if good else "❌ "
        print(f"{mark}Excel row {n:>2}  RAW:   {raw_txt}")
        if act["outcome"] == "loaded":
            flags = ", ".join(act["flags"]) or "-"
            print(f"              CLEAN: vendor={act['vendor']!r}  key={act['vendor_key']!r}  "
                  f"amount={act['amount']}  date={act['date']}  flags={flags}")
        elif act["outcome"] == "blank":
            print("              CLEAN: (blank row, ignored)")
        else:
            print(f"              {act['outcome'].upper()}: {act['reason']}")
        if expected and not agrees(expected[n], act):
            print(f"              EXPECTED: {expected[n]}")
        print()

    counts = {k: sum(a["outcome"] == k for a in actual.values()) for k in ("loaded", "skipped", "error")}
    flagged = sum(bool(cr.flags) for cr in cleaned.rows)
    print(f"Summary: {ws.max_row - 1} rows below header | loaded {counts['loaded']} | "
          f"skipped {counts['skipped']} | errors {counts['error']} | flagged {flagged}")
    if expected:
        print(f"Against labels: {ok} agree, {bad} differ")


if __name__ == "__main__":
    main()
